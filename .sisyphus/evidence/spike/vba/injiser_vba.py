#!/usr/bin/env python3
"""T3-spike: injiser vbaProject.bin i en openpyxl-generert .xlsx.

Steg:
1. Lag minimal .xlsx med openpyxl.
2. Åpne zip, patch [Content_Types].xml og xl/_rels/workbook.xml.rels,
   injiser xl/vbaProject.bin.
3. Lagre som .xlsm.
4. Verifiser med openpyxl (keep_vba=True) og Excel COM (VBComponents).

Inspirasjon: injiser_avkryssingsbokser() i lag_supportsystem.py.
"""

import os
import re
import sys
import zipfile
from openpyxl import Workbook, load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
EVIDENCE = os.path.join(BASE, "..", "..", "task-3-vba-spike.txt")
XLSX_INN = os.path.join(BASE, "spike_inn.xlsx")
XLSM_UT = os.path.join(BASE, "spike.xlsm")
BIN_INN = os.path.join(BASE, "vbaProject.bin")

CT_VBA = "application/vnd.ms-office.vbaProject"
CT_WORKBOOK_MACRO = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
CT_WORKBOOK = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
REL_VBA = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"


def logg(linje):
    print(linje)
    with open(EVIDENCE, "a", encoding="utf-8") as f:
        f.write(linje + "\n")


def lag_minimal_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Demo"
    ws["A1"] = "VBA-spike"
    ws["B2"] = 42
    wb.save(XLSX_INN)
    logg(f"Minimal .xlsx laget: {XLSX_INN}")


def neste_rid(rels_xml):
    rids = re.findall(r'Id="rId(\d+)"', rels_xml)
    return max((int(x) for x in rids), default=0) + 1


def injiser_vba():
    if not os.path.exists(BIN_INN):
        raise FileNotFoundError(f"Mangler {BIN_INN} -- kjør lag_donor.py først")

    with zipfile.ZipFile(XLSX_INN, "r") as zin:
        deler = {n: zin.read(n) for n in zin.namelist()}

    # Injiser binæren
    with open(BIN_INN, "rb") as f:
        deler["xl/vbaProject.bin"] = f.read()
    logg(f"vbaProject.bin injisert ({len(deler['xl/vbaProject.bin'])} bytes)")

    # Patch xl/_rels/workbook.xml.rels
    rels = deler["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid = f"rId{neste_rid(rels)}"
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{rid}" Type="{REL_VBA}" Target="vbaProject.bin"/>\n</Relationships>',
    )
    deler["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    logg(f"vbaProject-relasjon lagt til: {rid}")

    # Patch [Content_Types].xml
    ct = deler["[Content_Types].xml"].decode("utf-8")

    # Endre workbook content type til macroEnabled
    ct = re.sub(
        r'<Override[^>]*PartName="/xl/workbook\.xml"[^>]*/>',
        f'<Override PartName="/xl/workbook.xml" ContentType="{CT_WORKBOOK_MACRO}"/>',
        ct,
    )
    logg("Workbook ContentType endret til macroEnabled")

    # Legg til Override for vbaProject.bin hvis ikke finnes
    if CT_VBA not in ct:
        ct = ct.replace(
            "</Types>",
            f'<Override PartName="/xl/vbaProject.bin" ContentType="{CT_VBA}"/>\n</Types>',
        )
        logg("Override for vbaProject.bin lagt til")

    deler["[Content_Types].xml"] = ct.encode("utf-8")

    # Skriv ny zip
    tmp = XLSM_UT + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zut:
        for n, data in deler.items():
            zut.writestr(n, data)
    os.replace(tmp, XLSM_UT)
    logg(f"Makro-aktivert fil lagret: {XLSM_UT}")


def verifiser_openpyxl():
    wb = load_workbook(XLSM_UT, keep_vba=True)
    ok = "xl/vbaProject.bin" in wb.vba_archive.namelist()
    logg(f"openpyxl keep_vba sjekk: vbaProject.bin tilstede = {ok}")
    return ok


def verifiser_excel():
    import win32com.client as win32

    app = win32.Dispatch("Excel.Application")
    wb = None
    funnet = False
    try:
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Open(os.path.abspath(XLSM_UT))
        prosjekt = wb.VBProject
        navn = [c.Name for c in prosjekt.VBComponents]
        logg(f"VBComponents i spike.xlsm: {navn}")
        funnet = any("SayHello" in c.CodeModule.Lines(1, c.CodeModule.CountOfLines)
                     for c in prosjekt.VBComponents if c.Type == 1)
        logg(f"SayHello funnet i standardmodul = {funnet}")
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        app.Quit()
        logg("Excel avsluttet etter verifisering")
    return funnet


if __name__ == "__main__":
    os.makedirs(os.path.dirname(EVIDENCE), exist_ok=True)
    logg("=== injiser_vba.py start ===")
    lag_minimal_xlsx()
    injiser_vba()
    assert verifiser_openpyxl(), "openpyxl fant ikke vbaProject.bin"
    assert verifiser_excel(), "Excel fant ikke SayHello"
    logg("=== injiser_vba.py ferdig ===")
    logg("SPIKE OK: VBA-injeksjon + Excel-åpning + makro-presens verifisert")
