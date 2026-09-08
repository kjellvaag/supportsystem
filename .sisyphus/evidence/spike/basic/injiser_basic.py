# -*- coding: utf-8 -*-
"""
Injiserer StarBasic-modul(er) i en eksisterende .ods-fil.

Bruksmønster (spike):
    python .sisyphus/evidence/spike/basic/injiser_basic.py min.ods med_makro.ods

Med egen .bas-kilde:
    python injiser_basic.py min.ods med_makro.ods --bas basic/ModSok.bas --modul ModSok

Med verifisering (soffice headless round-trip xlsx):
    python injiser_basic.py min.ods med_makro.ods --verify

soffice-sti:
    - CLI:  --soffice "C:\Program Files\LibreOffice\program\soffice.exe"
    - Env:  SOFFICE=C:\Program Files\LibreOffice\program\soffice.exe
    - Default: C:\Program Files\LibreOffice\program\soffice.exe
"""
import argparse
import os
import subprocess
import sys
import tempfile
import zipfile

DEFAULT_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

HELLO_BAS = '''Sub Hello()
    MsgBox "hei"
End Sub
'''


def library_container_xml(bib_navn="Standard"):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:library-container PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "script.dtd">
<script:library-container xmlns:script="http://openoffice.org/2000/script">
  <script:library script:name="{bib_navn}" script:language="StarBasic"/>
</script:library-container>'''.encode("utf-8")


def library_index_xml(bib_navn, modul_navn):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "script.dtd">
<script:library xmlns:script="http://openoffice.org/2000/script" script:name="{bib_navn}" script:language="StarBasic" script:readonly="false">
  <script:module script:name="{modul_navn}" script:language="StarBasic"/>
</script:library>'''.encode("utf-8")


def module_xml(modul_navn, kildekode):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "script.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="{modul_navn}" script:language="StarBasic">
<script:source encoding="UTF-8"><![CDATA[{kildekode}]]></script:source>
</script:module>'''.encode("utf-8")


def patch_manifest(manifest_xml, bib_navn, modul_navn):
    """Legger til manifest:file-entry for Basic-mappen, -lc, -lb og modulen."""
    entries = f'''
 <manifest:file-entry manifest:full-path="Basic/" manifest:media-type="application/binary"/>
 <manifest:file-entry manifest:full-path="Basic/script-lc.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="Basic/{bib_navn}/" manifest:media-type="application/binary"/>
 <manifest:file-entry manifest:full-path="Basic/{bib_navn}/script-lb.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="Basic/{bib_navn}/{modul_navn}.xml" manifest:media-type="text/xml"/>'''
    return manifest_xml.replace("</manifest:manifest>", entries + "\n</manifest:manifest>")


def injiser(input_ods, output_ods, kildekode, modul_navn="Module1", bib_navn="Standard"):
    """Zip-rewrite injeksjon av én Basic-modul i en .ods."""
    if not os.path.exists(input_ods):
        raise FileNotFoundError(f"Inndatafil finnes ikke: {input_ods}")

    with zipfile.ZipFile(input_ods, "r") as zin:
        deler = {n: zin.read(n) for n in zin.namelist()}

    deler["Basic/script-lc.xml"] = library_container_xml(bib_navn)
    deler[f"Basic/{bib_navn}/script-lb.xml"] = library_index_xml(bib_navn, modul_navn)
    deler[f"Basic/{bib_navn}/{modul_navn}.xml"] = module_xml(modul_navn, kildekode)

    manifest = deler["META-INF/manifest.xml"].decode("utf-8")
    manifest = patch_manifest(manifest, bib_navn, modul_navn)
    deler["META-INF/manifest.xml"] = manifest.encode("utf-8")

    tmp = output_ods + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zut:
        for n, data in deler.items():
            zut.writestr(n, data)
    os.replace(tmp, output_ods)
    return output_ods


def verifiser_med_soffice(ods_path, soffice):
    """Sjekker at .ods tåler en headless ODS -> XLSX-runde med soffice."""
    if not os.path.exists(soffice):
        raise FileNotFoundError(f"soffice finnes ikke: {soffice}")

    with tempfile.TemporaryDirectory() as tmp:
        cmd = [
            soffice,
            "--headless",
            "--convert-to", "xlsx",
            "--outdir", tmp,
            ods_path,
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode != 0:
            raise RuntimeError(f"soffice konvertering feilet:\n{result.stdout}\n{result.stderr}")

        base = os.path.splitext(os.path.basename(ods_path))[0] + ".xlsx"
        xlsx_path = os.path.join(tmp, base)
        if not os.path.exists(xlsx_path):
            raise RuntimeError(f"soffice returnerte OK, men produserte ikke {base}")

        # Åpner med openpyxl for å bevise at filen er gyldig xlsx
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path)
        return wb.sheetnames


def main():
    parser = argparse.ArgumentParser(description="Injiser StarBasic i .ods")
    parser.add_argument("input_ods", help="Kilde .ods (uten makro)")
    parser.add_argument("output_ods", help="Mål .ods (med makro)")
    parser.add_argument("--bas", help="Sti til .bas-fil med kilden")
    parser.add_argument("--modul", default="Module1", help="Modulnavn (default Module1)")
    parser.add_argument("--bib", default="Standard", help="Biblioteknavn (default Standard)")
    parser.add_argument("--verify", action="store_true", help="Kjør soffice headless round-trip etter injeksjon")
    parser.add_argument("--soffice", default=None, help="Full sti til soffice.exe")
    args = parser.parse_args()

    if args.bas:
        with open(args.bas, "r", encoding="utf-8") as f:
            kildekode = f.read()
    else:
        kildekode = HELLO_BAS

    injiser(args.input_ods, args.output_ods, kildekode, args.modul, args.bib)
    print(f"Injisert {args.modul} i {args.output_ods}")

    if args.verify:
        soffice = args.soffice or os.environ.get("SOFFICE") or DEFAULT_SOFFICE
        ark = verifiser_med_soffice(args.output_ods, soffice)
        print(f"soffice round-trip OK: ark = {ark}")


if __name__ == "__main__":
    main()
