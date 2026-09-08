# -*- coding: utf-8 -*-
"""T11-verifisering: --target calc + fallback/excel + pytest."""
import datetime
import os
import re
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SOFFICE = os.environ.get("SOFFICE", r"C:\Program Files\LibreOffice\program\soffice.exe")
EVIDENCE = os.path.join(ROOT, ".sisyphus", "evidence", "task-11-ods.txt")


def log(msg):
    print(msg)
    with open(EVIDENCE, "a", encoding="utf-8") as f:
        f.write(msg + "\n")


def run(cmd, **kw):
    return subprocess.run(cmd, capture_output=True, text=True, **kw)


def normalize_core(xml):
    xml = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", r"\g<1>FIXED\g<2>", xml)
    xml = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", r"\g<1>FIXED\g<2>", xml)
    return xml


def content_stable(path, head_path):
    with zipfile.ZipFile(path, "r") as za, zipfile.ZipFile(head_path, "r") as zb:
        na, nb = sorted(za.namelist()), sorted(zb.namelist())
        if na != nb:
            return False, f"namelist diff: {set(na)^set(nb)}"
        for n in na:
            da, db = za.read(n), zb.read(n)
            if n == "docProps/core.xml":
                da = normalize_core(da.decode("utf-8")).encode("utf-8")
                db = normalize_core(db.decode("utf-8")).encode("utf-8")
            if da != db:
                return False, f"entry diff: {n}"
    return True, "OK"


def main():
    os.makedirs(os.path.dirname(EVIDENCE), exist_ok=True)
    with open(EVIDENCE, "w", encoding="utf-8") as f:
        f.write(f"T11-verifisering kjørt {datetime.datetime.now().isoformat()}\n")
        f.write("=" * 60 + "\n\n")

    log("1. --target calc")
    r = run([sys.executable, "lag_supportsystem.py", "--target", "calc"], cwd=ROOT)
    log(f"   exit={r.returncode}")
    if r.returncode != 0:
        log(f"   stdout: {r.stdout}\n   stderr: {r.stderr}")
        return 1
    log("   Supportsystem.ods generert")

    log("\n2. ODS Basic-parts + manifest")
    ods_path = os.path.join(ROOT, "Supportsystem.ods")
    with zipfile.ZipFile(ods_path, "r") as z:
        names = z.namelist()
        required = ["Basic/script-lc.xml", "Basic/Standard/script-lb.xml"]
        required += [f"Basic/Standard/Module{i}.xml" for i in range(1, 5)]
        for req in required:
            ok = req in names
            log(f"   {req}: {'OK' if ok else 'MANGLER'}")
            if not ok:
                return 1
        manifest = z.read("META-INF/manifest.xml").decode("utf-8")
        for entry in ("Basic/", "Basic/script-lc.xml", "Basic/Standard/", "Basic/Standard/script-lb.xml"):
            ok = f'manifest:full-path="{entry}"' in manifest
            log(f"   manifest {entry}: {'OK' if ok else 'MANGLER'}")
            if not ok:
                return 1

    log("\n3. ODS round-trip (soffice -> xlsx)")
    with tempfile.TemporaryDirectory(prefix="supportsystem_t11_rt_") as tmp:
        cmd = [SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", tmp, ods_path]
        r = run(cmd)
        log(f"   soffice exit={r.returncode}")
        if r.returncode != 0:
            log(f"   stdout: {r.stdout}\n   stderr: {r.stderr}")
            return 1
        xlsx_rt = os.path.join(tmp, "Supportsystem.xlsx")
        if not os.path.exists(xlsx_rt):
            log("   round-trip xlsx ble ikke produsert")
            return 1
        wb = load_workbook(xlsx_rt)
        log(f"   round-trip ark: {wb.sheetnames}")
        if len(wb.sheetnames) != 7:
            log("   FEIL: forventet 7 ark")
            return 1
        log("   round-trip OK (7 ark)")

    log("\n4. Fallback .xlsx")
    r = run([sys.executable, "lag_supportsystem.py"], cwd=ROOT)
    log(f"   exit={r.returncode}")
    if r.returncode != 0:
        return 1
    r = subprocess.run(["git", "show", "HEAD:Supportsystem.xlsx"], cwd=ROOT, capture_output=True)
    head_tmp = os.path.join(ROOT, "__head_xlsx__.tmp")
    with open(head_tmp, "wb") as f:
        f.write(r.stdout)
    ok, reason = content_stable(os.path.join(ROOT, "Supportsystem.xlsx"), head_tmp)
    os.remove(head_tmp)
    log(f"   innholdsstabilitet: {reason}")
    if not ok:
        return 1

    log("\n5. Excel .xlsm")
    r = run([sys.executable, "lag_supportsystem.py", "--target", "excel"], cwd=ROOT)
    log(f"   exit={r.returncode}")
    if r.returncode != 0:
        log(f"   stdout: {r.stdout}\n   stderr: {r.stderr}")
        return 1
    xlsm_path = os.path.join(ROOT, "Supportsystem.xlsm")
    with zipfile.ZipFile(xlsm_path, "r") as z:
        has_vba = "xl/vbaProject.bin" in z.namelist()
        ct = z.read("[Content_Types].xml").decode("utf-8")
    log(f"   vbaProject.bin: {'OK' if has_vba else 'MANGLER'}")
    log(f"   macroEnabled i [Content_Types].xml: {'OK' if 'macroEnabled' in ct else 'MANGLER'}")
    log(f"   vbaProject i [Content_Types].xml: {'OK' if 'vbaProject' in ct else 'MANGLER'}")

    log("\n6. pytest test_struktur.py")
    r = run([sys.executable, "-m", "pytest", "test_struktur.py", "-v"], cwd=ROOT)
    log(f"   exit={r.returncode}")
    for line in r.stdout.splitlines():
        log(f"   {line}")
    if r.returncode != 0:
        for line in r.stderr.splitlines():
            log(f"   ERR {line}")
        return 1

    log("\nT11 VERIFISERT OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
