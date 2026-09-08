# -*- coding: utf-8 -*-
"""
F3 manuell QA: strukturelle sjekker av fallback .xlsx, Excel .xlsm og Calc .ods.
Skriver evidens-filer til .sisyphus/evidence/final-qa/.
"""
import os
import sys
import zipfile
import json
import subprocess
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, ROOT)
os.chdir(ROOT)
from kontrakt import normalize

EV = r".sisyphus\evidence\final-qa"
os.makedirs(EV, exist_ok=True)

lines = []
def log(msg):
    print(msg)
    lines.append(msg)

def run(cmd):
    return subprocess.run(cmd, shell=True, capture_output=True, text=True, encoding="utf-8", errors="replace")

def sjekk_ren_tom(wb):
    saker = wb["Saker"]
    kb = wb["Kunnskapsbase"]
    count_saker = sum(1 for r in range(2, 502) if saker.cell(r, 3).value)
    count_kb = sum(1 for r in range(2, 302) if kb.cell(r, 2).value)
    return count_saker, count_kb

def har_norsk(ws, col, maxrow=12):
    for r in range(2, maxrow):
        v = ws.cell(r, col).value
        if v and any(c in str(v) for c in "æøåÆØÅ"):
            return True
    return False

# ============================================================
# Forbered REN tilstand for strukturelle sjekker
# ============================================================
log("=== FORBEREDELSE: ren tilstand for alle tre targets ===")
for f in ["Supportsystem.xlsx", "Supportsystem.xlsm", "Supportsystem.ods"]:
    if os.path.exists(f):
        os.remove(f)
run("python lag_supportsystem.py ren")
run("python lag_supportsystem.py --target excel ren")
run("python lag_supportsystem.py --target calc ren")

# ============================================================
# SCENARIO 1: Fallback .xlsx
# ============================================================
log("\n=== SCENARIO 1: Fallback .xlsx ===")
wb = load_workbook("Supportsystem.xlsx", data_only=False)
sheet_names = wb.sheetnames
expected = ["Dashboard", "Saker", "Søk KB", "Søk saker", "Kunnskapsbase", "Oppsett", "Instruksjoner"]
ok_sheets = sheet_names == expected
log(f"Ark-rekkefolge: {sheet_names} -> {'OK' if ok_sheets else 'FAIL'}")

codenames = {ws.title: ws.sheet_properties.codeName for ws in wb.worksheets}
ok_cn = all(v is None for v in codenames.values())
log(f"codeName (fallback skal vaere None): {codenames} -> {'OK' if ok_cn else 'FAIL'}")

saker = wb["Saker"]
b2 = saker["B2"].value
ok_b2 = isinstance(b2, str) and "NOW()" in b2.upper() and "$C" in b2 and "$H" in b2
log(f"Saker!B2 sirkulaer formel: {repr(b2)} -> {'OK' if ok_b2 else 'FAIL'}")

y2 = saker["Y2"].value
l2 = wb["Kunnskapsbase"]["L2"].value
ok_helpers = isinstance(y2, str) and y2.startswith("=") and isinstance(l2, str) and l2.startswith("=")
log(f"Hjelpekolonner Saker!Y2 har formel: {isinstance(y2, str) and y2.startswith('=')}, KB!L2 har formel: {isinstance(l2, str) and l2.startswith('=')} -> {'OK' if ok_helpers else 'FAIL'}")

count_saker, count_kb = sjekk_ren_tom(wb)
ok_empty = count_saker == 0 and count_kb == 0
log(f"Ren modus: saker={count_saker}, KB={count_kb} -> {'OK' if ok_empty else 'FAIL'}")

result_xlsx = ok_sheets and ok_cn and ok_b2 and ok_helpers and ok_empty
log(f"SCENARIO 1 RESULTAT: {'PASS' if result_xlsx else 'FAIL'}")

# ============================================================
# SCENARIO 2: Excel .xlsm
# ============================================================
log("\n=== SCENARIO 2: Excel .xlsm ===")
wbm = load_workbook("Supportsystem.xlsm", keep_vba=True)
has_vba = "xl/vbaProject.bin" in wbm.vba_archive.namelist()
log(f"vbaProject.bin finnes: {has_vba} -> {'OK' if has_vba else 'FAIL'}")

with zipfile.ZipFile("Supportsystem.xlsm") as z:
    ct = z.read("[Content_Types].xml").decode("utf-8", errors="ignore")
ok_ct = "macroEnabled" in ct and "vbaProject" in ct
log(f"[Content_Types].xml macroEnabled+vbaProject: {ok_ct} -> {'OK' if ok_ct else 'FAIL'}")

codenames_xlsm = {ws.title: ws.sheet_properties.codeName for ws in wbm.worksheets}
expected_cn = {
    "Dashboard": "Sheet1", "Saker": "Sheet2", "Søk KB": "Sheet3",
    "Søk saker": "Sheet4", "Kunnskapsbase": "Sheet5",
    "Oppsett": "Sheet6", "Instruksjoner": "Sheet7"
}
ok_cn_xlsm = codenames_xlsm == expected_cn
log(f"codeName mapping: {codenames_xlsm} -> {'OK' if ok_cn_xlsm else 'FAIL'}")

b2_xlsm = wbm["Saker"]["B2"].value
ok_b2_xlsm = not isinstance(b2_xlsm, str) or "=" not in b2_xlsm
log(f"Saker!B2 ikke formel: {repr(b2_xlsm)} -> {'OK' if ok_b2_xlsm else 'FAIL'}")

count_saker_m, count_kb_m = sjekk_ren_tom(wbm)
ok_empty_m = count_saker_m == 0 and count_kb_m == 0
log(f"Ren .xlsm: saker={count_saker_m}, KB={count_kb_m} -> {'OK' if ok_empty_m else 'FAIL'}")

result_xlsm = has_vba and ok_ct and ok_cn_xlsm and ok_b2_xlsm and ok_empty_m
log(f"SCENARIO 2 RESULTAT: {'PASS' if result_xlsm else 'FAIL'}")

# ============================================================
# SCENARIO 3: Calc .ods
# ============================================================
log("\n=== SCENARIO 3: Calc .ods ===")
with zipfile.ZipFile("Supportsystem.ods") as z:
    names = z.namelist()
    required = ["Basic/script-lc.xml", "Basic/Standard/script-lb.xml"]
    modules = [f"Basic/Standard/Module{i}.xml" for i in range(1, 5)]
    present = {n: n in names for n in required + modules}
    manifest = z.read("META-INF/manifest.xml").decode("utf-8", errors="ignore")

ok_basic = all(present.values())
log(f"Basic-moduler tilstede: {present} -> {'OK' if ok_basic else 'FAIL'}")

manifest_ok = all(x in manifest for x in ["Basic/", "script-lc.xml", "script-lb.xml"] + [f"Module{i}.xml" for i in range(1,5)])
log(f"Manifest-entries for Basic: {manifest_ok} -> {'OK' if manifest_ok else 'FAIL'}")

result_ods = ok_basic and manifest_ok
log(f"SCENARIO 3 RESULTAT: {'PASS' if result_ods else 'FAIL'}")

# ============================================================
# SCENARIO 4: Backup endelsesbevaring
# ============================================================
log("\n=== SCENARIO 4: Backup endelsesbevaring ===")
with open("backup_config.json", "r", encoding="utf-8") as f:
    original_cfg = f.read()

def test_backup(ext):
    cfg = json.loads(original_cfg)
    cfg["kildefil"] = f"Supportsystem{ext}"
    cfg["behold_antall"] = 2
    cfg["lokal_backupmappe"] = "backups"
    cfg["nettverk_backupmappe"] = ""
    with open("backup_config.json", "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)
    r = run(f"python backup.py")
    with open(os.path.join(EV, f"backup{ext}-output.txt"), "w", encoding="utf-8") as f:
        f.write((r.stdout or "") + "\n" + (r.stderr or "") + "\nexitcode=" + str(r.returncode))
    files = sorted([x for x in os.listdir("backups") if x.startswith("Supportsystem_") and x.endswith(ext)], reverse=True)
    ok = len(files) > 0 and all(x.endswith(ext) for x in files)
    log(f"Backup {ext}: {len(files)} fil(er), nyeste={files[0] if files else 'INGEN'} -> {'OK' if ok else 'FAIL'}")
    return ok

ok_bak_xlsm = test_backup(".xlsm")
ok_bak_ods = test_backup(".ods")
with open("backup_config.json", "w", encoding="utf-8") as f:
    f.write(original_cfg)
log("backup_config.json gjenopprettet")
result_backup = ok_bak_xlsm and ok_bak_ods
log(f"SCENARIO 4 RESULTAT: {'PASS' if result_backup else 'FAIL'}")

# ============================================================
# SCENARIO 5: Edge-cases
# ============================================================
log("\n=== SCENARIO 5: Edge-cases ===")

# æøå: generer sample-data-varianter
log("Genererer sample-data-varianter for norske tegn-sjekk...")
run("python lag_supportsystem.py")
run("python lag_supportsystem.py --target excel")
run("python lag_supportsystem.py --target calc")

wb_sample = load_workbook("Supportsystem.xlsx", data_only=True)
wb_sample_xlsm = load_workbook("Supportsystem.xlsm", data_only=True, keep_vba=True)

norsk_xlsx = har_norsk(wb_sample["Saker"], 3) or har_norsk(wb_sample["Kunnskapsbase"], 2)
norsk_xlsm = har_norsk(wb_sample_xlsm["Saker"], 3) or har_norsk(wb_sample_xlsm["Kunnskapsbase"], 2)
log(f"Norske tegn overlever .xlsx: {norsk_xlsx}, .xlsm: {norsk_xlsm}")

# For .ods: konverter med soffice og les tilbake
conv_dir = os.path.join(EV, "ods_conv")
os.makedirs(conv_dir, exist_ok=True)
for f in os.listdir(conv_dir):
    os.remove(os.path.join(conv_dir, f))
r = run(r'"C:\Program Files\LibreOffice\program\soffice.exe" --headless --convert-to xlsx --outdir "' + conv_dir + r'" Supportsystem.ods')
with open(os.path.join(EV, "ods-conv-output.txt"), "w", encoding="utf-8") as f:
    f.write((r.stdout or "") + "\n" + (r.stderr or "") + "\nexitcode=" + str(r.returncode))
try:
    wb_ods_conv = load_workbook(os.path.join(conv_dir, "Supportsystem.xlsx"), data_only=True)
    norsk_ods = har_norsk(wb_ods_conv["Saker"], 3) or har_norsk(wb_ods_conv["Kunnskapsbase"], 2)
except Exception as e:
    norsk_ods = False
    log(f"ODS konvertering/lesing feilet: {e}")
log(f"Norske tegn overlever .ods (etter soffice-konvertering): {norsk_ods}")

# Bindestrek-normalisering
norm_ok = normalize("Wi-Fi nede") == "wifi nede"
log(f"kontrakt.normalize('Wi-Fi nede') == 'wifi nede': {norm_ok}")

# Sjekk VBA/Basic kilde bruker Replace("-","") ikke Replace("-"," ")
with open("vba/ModFelles.bas", "r", encoding="cp1252", errors="replace") as f:
    vba_src = f.read()
with open("basic/ModFelles.bas", "r", encoding="utf-8", errors="replace") as f:
    basic_src = f.read()
vba_dash = 'Replace(s, "-", "")' in vba_src or 'Replace(t, "-", "")' in vba_src
basic_dash = 'Replace(s, "-", "")' in basic_src or 'Replace(t, "-", "")' in basic_src
log(f"VBA fjerner bindestrek (Replace(...,\"-\",\"\")): {vba_dash}")
log(f"Basic fjerner bindestrek (Replace(...,\"-\",\"\")): {basic_dash}")

# Voksende lister: sjekk at definerte navn finnes (strukturelt)
names = wb.defined_names
list_names = ["Kanaler", "Kategorier", "Statuser", "Prioriteter", "Ansatte"]
ok_names = all(n in names for n in list_names)
log(f"Dynamiske navn for voksende lister: {ok_names} (fantes: {list(names.keys())})")

result_edge = (norsk_xlsx and norsk_xlsm and norsk_ods and norm_ok and
               vba_dash and basic_dash and ok_names)
log(f"SCENARIO 5 RESULTAT: {'PASS' if result_edge else 'FAIL'}")

# ============================================================
# SCENARIO 6: Kryss-integrasjon (alle tre targets etter hverandre)
# ============================================================
log("\n=== SCENARIO 6: Kryss-integrasjon ===")
for f in ["Supportsystem.xlsx", "Supportsystem.xlsm", "Supportsystem.ods"]:
    if os.path.exists(f):
        os.remove(f)
log("Slettet eksisterende Supportsystem.*")
res = [
    run("python lag_supportsystem.py"),
    run("python lag_supportsystem.py --target excel"),
    run("python lag_supportsystem.py --target calc")
]
for i, r in enumerate(res, 1):
    target = ["xlsx", "xlsm", "ods"][i-1]
    log(f"Generer {target}: exit={r.returncode}")
exists = {f: os.path.exists(f) for f in ["Supportsystem.xlsx", "Supportsystem.xlsm", "Supportsystem.ods"]}
log(f"Alle tre filer eksisterer: {exists}")
result_cross = all(r.returncode == 0 for r in res) and all(exists.values())
log(f"SCENARIO 6 RESULTAT: {'PASS' if result_cross else 'FAIL'}")

# ============================================================
# SCENARIO 7: test_struktur.py og valider_soffice.py (henvisning)
# ============================================================
log("\n=== SCENARIO 7: Automatiske valideringer ===")
pytest_path = os.path.join(EV, "pytest-output.txt")
soffice_path = os.path.join(EV, "soffice-output.txt")
log(f"Leter etter evidensfiler: {pytest_path}, {soffice_path}")

try:
    with open(pytest_path, "r", encoding="utf-8", errors="replace") as f:
        pytest_out = f.read()
    pytest_ok = "17 passed" in pytest_out
    log(f"pytest-output.txt har '17 passed': {pytest_ok}")
except Exception as e:
    pytest_ok = False
    log(f"Kunne ikke lese pytest-output.txt: {e}")
log(f"test_struktur.py 17 passed: {pytest_ok}")

try:
    with open(soffice_path, "r", encoding="utf-8", errors="replace") as f:
        soffice_out = f.read()
    soffice_ok = "PASS (alle 2 round-trip-konverteringer lyktes)" in soffice_out
    log(f"soffice-output.txt har PASS: {soffice_ok}")
except Exception as e:
    soffice_ok = False
    log(f"Kunne ikke lese soffice-output.txt: {e}")
log(f"valider_soffice.py PASS: {soffice_ok}")

# ============================================================
# Oppsummering
# ============================================================
scenarios = [result_xlsx, result_xlsm, result_ods, result_backup, result_edge, result_cross, pytest_ok, soffice_ok]
passed = sum(scenarios)
total = len(scenarios)
log("\n" + "=" * 60)
log(f"OPPSUMMERING: Scenarios {passed}/{total} PASS")
log(f"  1 fallback .xlsx : {'PASS' if result_xlsx else 'FAIL'}")
log(f"  2 Excel .xlsm    : {'PASS' if result_xlsm else 'FAIL'}")
log(f"  3 Calc .ods      : {'PASS' if result_ods else 'FAIL'}")
log(f"  4 Backup         : {'PASS' if result_backup else 'FAIL'}")
log(f"  5 Edge-cases     : {'PASS' if result_edge else 'FAIL'}")
log(f"  6 Kryss-integr.  : {'PASS' if result_cross else 'FAIL'}")
log(f"  7 test_struktur  : {'PASS' if pytest_ok else 'FAIL'}")
log(f"  8 valider_soffice: {'PASS' if soffice_ok else 'FAIL'}")
log("=" * 60)

with open(os.path.join(EV, "f3-struktur-sjekk.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

exit(0 if passed == total else 1)
