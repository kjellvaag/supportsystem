import zipfile, os, re, subprocess, tempfile, shutil
from openpyxl import load_workbook

evidence_dir = ".sisyphus/evidence/final-qa"
os.makedirs(evidence_dir, exist_ok=True)

# 1. Compare generated .xlsx with git HEAD version
print("=== Scenario 1: Fallback content-stability ===")
print()

# Get git HEAD version of Supportsystem.xlsx
head_bytes = subprocess.check_output(["git", "show", "HEAD:Supportsystem.xlsx"])

with open("Supportsystem.xlsx", "rb") as f:
    gen_bytes = f.read()

print(f"Generated size: {len(gen_bytes)} bytes")
print(f"HEAD size: {len(head_bytes)} bytes")

# Normalize docProps timestamps in both zip files
def normalize_zip(zb):
    zin = zipfile.ZipFile(zb)
    out = {}
    for n in zin.namelist():
        data = zin.read(n)
        if n.startswith("docProps/") and n.endswith(".xml"):
            # Replace all date/time strings with a fixed placeholder
            data = re.sub(rb"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z?", b"NORMALIZED-TIMESTAMP", data)
        out[n] = data
    return out

gen_norm = normalize_zip(gen_bytes)
head_norm = normalize_zip(head_bytes)

# Compare
only_gen = sorted(set(gen_norm) - set(head_norm))
only_head = sorted(set(head_norm) - set(gen_norm))
both = sorted(set(gen_norm) & set(head_norm))

diffs = []
for n in both:
    if gen_norm[n] != head_norm[n]:
        diffs.append(n)

print(f"Entries only in generated: {only_gen}")
print(f"Entries only in HEAD: {only_head}")
print(f"Differing entries (after timestamp normalization): {diffs}")

# Write normalized zips for inspection
with zipfile.ZipFile(os.path.join(evidence_dir, "gen_normalized.xlsx.zip"), "w", zipfile.ZIP_DEFLATED) as z:
    for n, d in gen_norm.items():
        z.writestr(n, d)
with zipfile.ZipFile(os.path.join(evidence_dir, "head_normalized.xlsx.zip"), "w", zipfile.ZIP_DEFLATED) as z:
    for n, d in head_norm.items():
        z.writestr(n, d)

# 2. Verify circular reference in Saker!B12 (B2 overwritten by sample data)
wb = load_workbook("Supportsystem.xlsx")
saker = wb["Saker"]
print(f"\nSaker!B2 value: {saker['B2'].value!r}")
print(f"Saker!B12 value: {saker['B12'].value!r}")
expected_b12_actual = '=IF($C12&$H12="","",IF($B12="",NOW(),$B12))'
print(f"Expected B12 formula: {expected_b12_actual!r}")
print(f"B12 matches: {saker['B12'].value == expected_b12_actual}")

# 3. Run ren mode and verify empty
print("\n=== Running ren mode ===")
os.system("python lag_supportsystem.py ren")
wb_ren = load_workbook("Supportsystem.xlsx")
print(f"Ren mode Saker!B2: {wb_ren['Saker']['B2'].value!r}")
print(f"Ren mode Saker!C2: {wb_ren['Saker']['C2'].value!r}")
print(f"Ren mode KB B2: {wb_ren['Kunnskapsbase']['B2'].value!r}")

# Count sample data rows
sample_count = sum(1 for r in range(2, 12) if wb_ren['Saker'].cell(row=r, column=3).value is not None)
print(f"Sample data rows in ren mode: {sample_count}")

# Save results
with open(os.path.join(evidence_dir, "scenario1_result.txt"), "w", encoding="utf-8") as f:
    f.write("SCENARIO 1 RESULT\n")
    f.write(f"Generated size: {len(gen_bytes)} bytes\n")
    f.write(f"HEAD size: {len(head_bytes)} bytes\n")
    f.write(f"Entries only in generated: {only_gen}\n")
    f.write(f"Entries only in HEAD: {only_head}\n")
    f.write(f"Differing entries (after timestamp normalization): {diffs}\n")
    f.write(f"Saker!B12 formula matches: {saker['B12'].value == expected_b12_actual}\n")
    f.write(f"Ren mode sample data rows: {sample_count}\n")
    f.write(f"PASS: {len(diffs) == 0 and saker['B12'].value == expected_b12_actual and sample_count == 0}\n")
